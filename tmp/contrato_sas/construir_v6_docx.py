# -*- coding: utf-8 -*-
# Contrato de Socios TECNIFIN S.A.S. - version 6 (Word).
# Base: version 5 de Christian Cuenca (21-sep-2026) + acuerdos de la reunion del 21-sep (minutos 0-11)
# + correcciones de contexto. Lo cambiado va RESALTADO en amarillo; lo que pide decision de los socios va como
# COMENTARIO de Word (no se cambia el texto). Anexos A, B, C, E: Jorge y Franklin. Anexos D, F, H: Christian (esqueleto).
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"C:\Users\DELL\Documents\TRABAJOS ADICIONALES\GUTT COMPANY SAS\Contrato_Socios_TECNIFIN_SAS_v6.docx"
NAVY = RGBColor(0x1F, 0x38, 0x64)
AZUL = RGBColor(0x2F, 0x54, 0x96)
AUTOR, INI = "Jorge Tuquinga (revision v6)", "JT"

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2); s.left_margin = s.right_margin = Cm(2.3)
    hp = s.header.paragraphs[0]; hp.text = "TECNIFIN S.A.S. - CONTRATO DE SOCIOS Y COMPROMISO DE EJECUCIÓN"
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER; hp.runs[0].font.size = Pt(8.5); hp.runs[0].bold = True
    fp = s.footer.paragraphs[0]; fp.text = "Quito, septiembre de 2026 · Versión 6 (22-sep-2026)"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; fp.runs[0].font.size = Pt(8.5)

def resaltar(run): run.font.highlight_color = WD_COLOR_INDEX.YELLOW

def p(texto, cambio=False, negrita_inicial=None, comentario=None, centro=False, tam=None):
    par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(5)
    if centro: par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else: par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    runs = []
    if negrita_inicial:
        r = par.add_run(negrita_inicial + " "); r.bold = True; runs.append(r)
    r = par.add_run(texto); runs.append(r)
    for x in runs:
        if cambio: resaltar(x)
        if tam: x.font.size = Pt(tam)
    if comentario: doc.add_comment(runs, text=comentario, author=AUTOR, initials=INI)
    return par

def h(texto, nivel=1):
    par = doc.add_paragraph(); r = par.add_run(texto); r.bold = True
    r.font.color.rgb = NAVY if nivel == 1 else AZUL; r.font.size = Pt(13 if nivel == 1 else 11.5)
    par.paragraph_format.space_before = Pt(10 if nivel == 1 else 6); par.paragraph_format.keep_with_next = True
    return par

def tabla(filas, anchos=None, cambio_filas=(), encabezado=True, tam=9.5):
    t = doc.add_table(rows=len(filas), cols=len(filas[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(filas):
        for j, val in enumerate(fila):
            c = t.cell(i, j); c.text = ""
            r = c.paragraphs[0].add_run(str(val)); r.font.size = Pt(tam)
            if i == 0 and encabezado: r.bold = True
            if i in cambio_filas: resaltar(r)
            if anchos: c.width = Cm(anchos[j])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def salto(): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ------------------------------------------------------------------ PORTADA
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("CONTRATO DE SOCIOS Y COMPROMISO DE EJECUCIÓN"); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = NAVY
p("CONSTITUCIÓN DE TECNIFIN S.A.S. E IMPLEMENTACIÓN DE LA PLATAFORMA", centro=True).runs[0].bold = True
p("Proyecto: SISTEMA FINANCIERO", centro=True)
p("La redacción societaria, laboral, tributaria, de propiedad intelectual y protección de datos deberá ser validada por asesoría jurídica ecuatoriana antes de la firma.",
  comentario="Se retiraron de la portada «Documento de trabajo para negociación» y «Versión de revisión de Christian Cuenca», acordado el 19-sep (quitar referencias a borrador). Esta advertencia de validación jurídica se conserva: los socios deciden si la retiran para la versión de firma.")
p("Leyenda de esta versión: el texto resaltado en amarillo es lo que cambió respecto de la versión 5 (revisión de Christian Cuenca, 21-sep-2026). Los puntos que requieren decisión de los socios van como comentarios al margen y no alteran el texto. Antes de firmar se retiran el resaltado, los comentarios y esta leyenda.",
  cambio=True, tam=9.5)

# ------------------------------------------------------------------ 1
h("1. COMPARECIENTES")
p("En la ciudad de Quito, Distrito Metropolitano, a los ____ días del mes de __________ de 2026, comparecen libre y voluntariamente las personas que se detallan a continuación, en adelante LOS SOCIOS, quienes acuerdan celebrar el presente contrato preparatorio de socios y compromiso de ejecución:")
tabla([["Nombre", "Cédula", "Cargo / función", "Calidad"],
       ["Víctor Emilio Cuenca Caraguay", "1101475232", "Gerente", "Socio fundador / aportante"],
       ["Franklin Sebastián Lechón Quilo", "1753860053", "Presidente / Desarrollador", "Socio fundador / desarrollador"],
       ["Christian Cuenca", "1716860851", "Jefe de Proyecto y Responsable de Infraestructura Tecnológica y Seguridad de la Información", "Socio fundador / aportante técnico"],
       ["Jorge Vladimir Tuquinga Tituaña", "1720884012", "Desarrollador", "Socio fundador / desarrollador"],
       ["____________________ (inversionista)", "____________", "Inversionista", "Socio fundador / aportante"]],
      anchos=[5, 2.7, 5.3, 4], cambio_filas=(1, 2, 5))
p("Cédulas y nombres completos: verificar cada socio el suyo. La cédula de Franklin se tomó de su mensaje del 17-sep en el grupo de socios; el nombre completo de Christian falta. En la versión 5 las cédulas del bloque de firmas estaban desalineadas con los nombres.",
  cambio=True, tam=9, comentario="Víctor Cuenca completa el nombre y la cédula del inversionista el 23-sep (reunión del 21-sep, min 4:39-4:45).")

# ------------------------------------------------------------------ 2
h("2. ANTECEDENTES")
p("2.1. Los socios desarrolladores declaran haber construido la plataforma SISTEMA FINANCIERO EPS / GUTT SYSTEM, actualmente concebida para operación mono-tenant y respecto de la cual manifiestan un grado estimado de avance aproximado del sesenta por ciento (60%). Este porcentaje es exclusivamente declarativo y no constituye por sí mismo valoración económica, aceptación técnica, acreditación de titularidad ni porcentaje del valor de la futura sociedad.")
p("2.2. Los comparecientes han resuelto constituir una Sociedad por Acciones Simplificada bajo la denominación tentativa TECNIFIN S.A.S., con GUTT COMPANY S.A.S. como denominación alterna conforme a la cláusula 4.1, cuyo giro comprenderá el desarrollo, implementación, licenciamiento, soporte y comercialización de soluciones informáticas financieras, incluida su prestación como servicio en la nube.", cambio=True)
p("2.3. El proyecto contempla un período referencial de ejecución de diez (10) meses, contados desde el 1 de octubre de 2026 y con cierre previsto el 31 de julio de 2027, para completar, adaptar, documentar, desplegar y poner en producción la plataforma, incluyendo desarrollo, arquitectura multi-tenant, migración, infraestructura tecnológica, seguridad de la información, continuidad, documentación y capacitación.", cambio=True)
p("2.4. Víctor Cuenca prevé un aporte en numerario de hasta USD 20.000, sujeto a la matriz vinculante de desembolsos y a los hitos de aceptación. Adicionalmente, pondrá a disposición de la sociedad un espacio físico para alojamiento de infraestructura, en las condiciones establecidas en este contrato.")
p("2.5. El INVERSIONISTA prevé un aporte en numerario de hasta USD 20.000, sujeto a la matriz vinculante de desembolsos y a los hitos de aceptación.")
p("2.6. Christian Cuenca realizará un aporte especializado de trabajo en Jefatura de Proyecto, Infraestructura Tecnológica, Seguridad de la Información y Administración Técnica. Su valoración referencial inicial será documentada mediante registro de horas y entregables y no se confundirá con salario, pago automático en efectivo ni gasto de capital.",
  comentario="Decidir: «pago automático en efectivo» puede leerse en contra de los USD 4.970 remunerables de la cláusula 10.5. Propuesta: «…y la parte remunerable de la cláusula 10.5 se pagará contra trabajo documentado».")
p("2.7. Se prevé la incorporación de un quinto socio/inversionista, cuya identificación completa se encuentra pendiente de incorporar al presente instrumento. Su aporte previsto será el señalado en la cláusula 2.5, y su calendario de desembolsos, participación y demás condiciones de ingreso deberán quedar aprobados y documentados en los instrumentos correspondientes.")

# ------------------------------------------------------------------ 3
h("3. OBJETO DEL CONTRATO")
p("El presente instrumento tiene por objeto formalizar los compromisos para constituir TECNIFIN S.A.S. y ejecutar el proyecto, distinguiendo expresamente: (i) software y activos tecnológicos existentes; (ii) derechos de propiedad intelectual y su transferencia; (iii) desarrollo funcional y técnico pendiente; (iv) aportes en numerario; (v) trabajo de dirección, administración, infraestructura tecnológica y seguridad de la información; (vi) trabajo y entregables correspondientes a desarrollo; (vii) derecho de uso de espacio físico; y (viii) demás activos, servicios o derechos que se acuerden.")
p("Los anexos técnicos y económicos forman parte integrante e inseparable de este contrato y obligan a los socios desde su firma. Contienen, según corresponda, alcance, responsables, plazos, evidencia, criterios de aceptación, valoración y condiciones de entrega.", cambio=True,
  comentario="Reunión 21-sep (min 4:01-4:09, Víctor): «que quede claro que los anexos son parte del contrato y lo que firmamos nos obligamos a eso». Por eso se firman junto con el contrato.")

# ------------------------------------------------------------------ 4
h("4. DE LA SOCIEDAD, GOBIERNO Y FUNCIONES")
p("4.1. Denominación. La sociedad tendrá como denominación principal TECNIFIN S.A.S. y como denominación alterna GUTT COMPANY S.A.S. Ambas se presentarán, en ese orden, para la reserva de nombre y la constitución ante la autoridad competente; si la principal no estuviere disponible, se adoptará la alterna sin alterar las demás obligaciones.", cambio=True,
  comentario="Confirmado en la reunión del 21-sep (min 8:10-8:24): nombre principal TECNIFIN, segundo nombre GUTT COMPANY.")
p("4.2. Domicilio. Cantón Quito, provincia de Pichincha, República del Ecuador.")
p("4.3. Naturaleza preparatoria. Este contrato es un acuerdo interno de compromisos entre los socios, previo a la constitución de la sociedad; no forma parte de los estatutos. Rige desde su suscripción y servirá de base para la constitución y los instrumentos societarios definitivos. Las obligaciones de trabajo y de pago que contiene comienzan el 1 de octubre de 2026, aunque la legalización de la sociedad esté pendiente. Las estipulaciones que por su naturaleza deban incorporarse al estatuto, pactos de accionistas, instrumentos de cesión u otros documentos serán formalizadas oportunamente.", cambio=True)
h("4.4. Jefatura de Proyecto", 2)
p("Christian Cuenca, como JEFE DE PROYECTO, será responsable de la planificación, coordinación y seguimiento integral del cronograma; gestión de riesgos, dependencias, cambios, incidencias y actas; coordinación entre desarrollo, infraestructura, administración y clientes; seguimiento de entregables; organización de pruebas y validaciones; y emisión o coordinación de actas de aceptación conforme a criterios objetivos. La aceptación no se producirá por silencio.")
h("4.5. Infraestructura Tecnológica y Seguridad de la Información", 2)
p("Christian Cuenca, como RESPONSABLE DE INFRAESTRUCTURA TECNOLÓGICA Y SEGURIDAD DE LA INFORMACIÓN, será responsable de diseñar, documentar, implementar, configurar, administrar, monitorear y mantener la plataforma tecnológica sobre la cual operará el sistema, incluyendo cómputo, virtualización o contenedores cuando corresponda, sistemas operativos, almacenamiento, redes, conectividad, segmentación, firewall, VPN, proxy inverso, TLS/certificados, servicios de infraestructura, accesos administrativos, hardening, registros, monitoreo, alertamiento, respaldos, restauración, continuidad, recuperación ante desastres, capacidad, rendimiento, parchado y gestión de vulnerabilidades.")
p("Las decisiones sobre código, modelo de datos y lógica funcional serán responsabilidad primaria de los desarrolladores. Las decisiones de infraestructura y seguridad serán responsabilidad primaria del Responsable de Infraestructura y Seguridad. Las decisiones que involucren ambas capas deberán documentarse y aprobarse mediante Decisiones de Arquitectura (ADR) cuando sean relevantes.")
h("4.6. Desarrollo de Software", 2)
p("Franklin Lechón y Jorge Vladimir Tuquinga Tituaña, como RESPONSABLES DE DESARROLLO, serán responsables del análisis funcional y normativo (incluyendo SEPS, UAF y Catálogo Único de Cuentas cuando corresponda); arquitectura de aplicación y datos; diseño e implementación del modelo multi-tenant; desarrollo, adaptación y mantenimiento de backend, frontend y móvil; scripts versionados y migraciones de base de datos; reportería regulatoria; integraciones y APIs; pruebas automatizadas y de aceptación funcional; migración y reconciliación de datos; hardening, gestión de secretos y monitoreo de la capa de aplicación; validación lógica de la base de datos después de restauraciones; documentación funcional y técnica; manuales de usuario; capacitación funcional; y soporte de estabilización durante la puesta en producción. Ambos responderán frente a la sociedad por los entregables de desarrollo, sin perjuicio de la distribución interna de tareas registrada en el tablero del proyecto.")
h("4.7. Gerencia y aporte en numerario", 2)
p("Víctor Cuenca, como GERENTE, será responsable de: proveer y actualizar el flujo de caja de aportes del Anexo D; administrar el capital de trabajo y los fondos del proyecto; el tratamiento contable de los aportes, del software y de las cuentas por pagar durante la constitución de la sociedad; efectuar su aporte en numerario conforme a la cláusula 10.4 y a los hitos de la cláusula 8.4.1; y el espacio físico de la cláusula 12. Sus funciones y compromisos se detallan en el Anexo B.", cambio=True,
  comentario="Reunión 21-sep (min 2:07-2:21, Víctor): «dónde está el trabajo que voy a hacer yo y tiene que hacer el nuevo inversionista… eso debe estar detallado en el anexo». Texto de funciones tomado de la versión 3 y de la reunión del 19-sep; Víctor valida.")
h("4.8. Inversionista", 2)
p("El INVERSIONISTA efectuará su aporte en numerario conforme a la cláusula 2.5 y a los hitos de la cláusula 8.4.1, y apoyará la socialización, el relacionamiento institucional y la comercialización del producto. Sus funciones y compromisos se detallan en el Anexo B y se confirman con su incorporación.", cambio=True,
  comentario="Funciones propuestas en la reunión del 17-sep (socialización, cabildeo y comercialización). Víctor las confirma con el inversionista el 23-sep.")

# ------------------------------------------------------------------ 5
h("5. ALCANCE Y FASES DEL PROYECTO")
tabla([["Fase", "Objeto principal", "Duración referencial", "Condición de cierre"],
       ["1", "Arquitectura y núcleo multi-tenant", "2 meses", "Documento de arquitectura, ADR críticas aprobadas y núcleo validado"],
       ["2", "Adaptación de módulos", "3 meses", "Módulos acordados adaptados y pruebas sin discrepancias críticas"],
       ["3", "Migración e integración", "1 mes", "Migración controlada, reconciliación, integración E2E y rollback documentado"],
       ["4", "Infraestructura física/lógica", "1 mes", "Infraestructura instalada, aprovisionada e inventariada"],
       ["5", "Hardening, seguridad, observabilidad y continuidad", "1 mes", "Controles implementados, monitoreo y pruebas de restore/DR con evidencia"],
       ["6", "Documentación y capacitación", "1 mes", "Manuales, runbooks y capacitación acordada"],
       ["7", "Producción y entrega", "1 mes", "Puesta en producción y acta formal de cierre"]], anchos=[1.3, 5.3, 3, 7.4])
p("El plazo total referencial es de diez (10) meses. Las fechas de cada fase y sus entregables constan en el Anexo A. Los plazos podrán modificarse únicamente mediante control de cambios aprobado por los socios conforme a este contrato.", cambio=True,
  comentario="Decidir: la base de TECNIFIN nace en blanco (decisión del 20-sep). La Fase 3 se ejecuta como integración E2E sobre la base de demostración, y la migración solo aplica cuando una cooperativa traiga datos de su sistema. Así queda redactado en el Anexo A sin cambiar esta tabla.")

# ------------------------------------------------------------------ 6
h("6. ARQUITECTURA MULTI-TENANT Y DECISIONES TÉCNICAS")
p("6.1. Los desarrolladores deberán diseñar, implementar y documentar la arquitectura multi-tenant antes de la adaptación masiva de módulos. El documento deberá cubrir como mínimo: objetivo y alcance; estado actual; alternativas evaluadas; modelo seleccionado; identificación de tenant; aislamiento de datos; modelo de aplicación; modelo de base de datos; IAM; auditoría; configuración por tenant; backup/restore; migración; rendimiento; escalabilidad; observabilidad; seguridad; continuidad; impacto operativo; riesgos y deuda técnica.")
p("6.2. Las decisiones arquitectónicas relevantes deberán registrarse mediante ADR, incluyendo contexto, requisitos, opciones evaluadas, decisión, justificación, consecuencias, riesgos, pruebas requeridas, criterio de aceptación y aprobaciones.")
p("6.3. Ninguna ADR crítica se considerará aprobada por silencio. Las decisiones que afecten aislamiento de datos, seguridad, disponibilidad, respaldo, recuperación o capacidad requerirán revisión expresa de Desarrollo e Infraestructura/Seguridad antes de su implementación definitiva.")

# ------------------------------------------------------------------ 7
h("7. COMPROMISO DE EJECUCIÓN DE LOS DESARROLLADORES")
p("7.1. Franklin Lechón y Jorge Vladimir Tuquinga Tituaña, en calidad de socios desarrolladores, se obligan a ejecutar y entregar el desarrollo funcional y técnico pendiente dentro del alcance, fases, plazos y criterios de aceptación aprobados.")
p("7.2. Su obligación comprende el diseño y desarrollo de la aplicación, arquitectura de aplicación y datos, adaptación multi-tenant, migración lógica, pruebas de aplicación, documentación funcional y técnica, integraciones, hardening de la capa de aplicación, monitoreo de aplicación y validación lógica de base de datos después de restauraciones.")
p("7.3. La instalación, configuración, operación y seguridad de la infraestructura no se considerará obligación exclusiva de los desarrolladores. Estas actividades serán dirigidas por Christian Cuenca en su función de Infraestructura y Seguridad, con la coordinación técnica de Desarrollo cuando exista impacto en aplicación, base de datos o despliegue.")
p("7.4. Los desarrolladores deberán mantener operativo, dentro de las condiciones previamente existentes y técnicamente razonables, el sistema que ya se encuentre en producción durante la transición, salvo ventanas de mantenimiento o causas justificadas.")
p("7.5. Los retrasos imputables a los responsables de una fase deberán ser subsanados sin generar automáticamente pagos adicionales. No se considerará retraso imputable aquel causado por fuerza mayor, retrasos de terceros no controlables, falta de desembolso aprobado, indisponibilidad de insumos requeridos o cambios de alcance formalmente autorizados.")
p("7.6. Entregables de desarrollo por fase. Como mínimo, los socios desarrolladores deberán producir y mantener los siguientes entregables, con la misma exigencia de documentación y aceptación prevista para infraestructura y seguridad:")
tabla([["ID", "Entregable mínimo", "Fase", "Responsable primario"],
       ["ARQ-01", "Documento de arquitectura multi-tenant", "1", "Desarrollo"],
       ["ARQ-02", "ADR críticas: tenant, aislamiento y modelo de datos", "1", "Desarrollo + Christian según capa"],
       ["DAT-01", "Modelo de datos multi-tenant y scripts versionados", "1", "Desarrollo"],
       ["APP-01", "Núcleo multi-tenant validado: tenant, IAM, auditoría y configuración", "1", "Desarrollo"],
       ["MOD-01", "Adaptación: socios, cuentas y caja", "2", "Desarrollo"],
       ["MOD-02", "Adaptación: créditos, cartera y plazo fijo", "2", "Desarrollo"],
       ["MOD-03", "Contabilidad y reportería regulatoria SEPS/UAF", "2", "Desarrollo"],
       ["QA-01", "Pruebas automatizadas y reconciliación sin discrepancias críticas", "2-3", "Desarrollo"],
       ["MIG-01", "Migración con reconciliación y rollback documentado", "3", "Desarrollo"],
       ["INT-01", "Integraciones, APIs y pruebas E2E", "3", "Desarrollo"],
       ["DEP-01", "Paquete de despliegue y soporte a instalación", "4", "Desarrollo + Christian"],
       ["APP-02", "Hardening, secretos y monitoreo de aplicación", "5", "Compartido según capa"],
       ["BKP-03", "Validación lógica de BD después de restauraciones", "5", "Desarrollo + Christian"],
       ["DOC-02", "Documentación técnica y manuales", "6", "Desarrollo"],
       ["CAP-02", "Capacitación funcional", "6", "Desarrollo"],
       ["PRD-01", "Puesta en producción y estabilización", "7", "Desarrollo + Christian"]], anchos=[1.8, 8.8, 1.3, 5.1])
p("Nota de edición: en la versión 5 esta tabla se desalineó al exportar a PDF (las descripciones no coincidían con los códigos). Se reconstruyó con la correspondencia de la versión 3.", cambio=True, tam=9)

# ------------------------------------------------------------------ 8
h("8. ACEPTACIÓN, CONTROL DE CAMBIOS Y DESEMBOLSOS")
p("8.1. Todo entregable deberá tener evidencia objetiva y criterios de aceptación previamente definidos. Concluida una fase, el responsable notificará formalmente al Jefe de Proyecto, quien coordinará la revisión técnica y emitirá acta de aceptación u observaciones.")
p("8.2. Ningún entregable crítico se aceptará por silencio. La falta de pronunciamiento no equivaldrá a aceptación, renuncia a observaciones ni conformidad técnica.")
p("8.3. Las observaciones fundadas deberán subsanarse dentro del plazo acordado para el entregable o, a falta de plazo específico, dentro de diez (10) días hábiles, salvo que por complejidad técnica se apruebe otro plazo.")
p("8.4. Cada desembolso quedará vinculado a un hito verificable y a su aceptación. La matriz de desembolsos deberá identificar: hito, entregable, evidencia, responsable, monto, fecha, condición de pago, retención si existiere y tratamiento de observaciones.")
p("8.4.1. Cronograma de hitos como guía para desembolsos. Los aportes en numerario no se entenderán exigibles de forma automática por el mero transcurso de cada mes. Su desembolso se programará conforme a necesidades reales y aprobadas del proyecto, tomando como guía los hitos siguientes y el flujo de caja de aportes del Anexo D. Todo desembolso deberá conservar trazabilidad, soporte y relación con el hito correspondiente. La fecha efectiva de compra o pago de hardware podrá ajustarse a las condiciones comerciales del proveedor mediante aprobación documentada, sin alterar el límite de aporte comprometido por cada socio.")
tabla([["Mes", "Hito", "Orientación para desembolsos"],
       ["Mes 1 (oct-2026)", "Fase 1 - arranque de la finalización y validación del esquema multi-tenant. Adquisición de dominio y hosting (tope conjunto USD 150).", "Liberar únicamente gastos y compensaciones aprobadas para el arranque."],
       ["Mes 2 (nov-2026)", "Fase 1 - cierre: esquema validado, íntegro y con auditoría cubierta.", "Desembolso sujeto a evidencia del cierre de Fase 1 y obligaciones aprobadas."],
       ["Mes 3 (dic-2026)", "Fase 2 - adaptación de módulos de créditos y contabilidad.", "Financiar actividades aprobadas de Fase 2."],
       ["Mes 4 (ene-2027)", "Fase 2 - adaptación de reportería SEPS y proceso de cartera.", "Financiar actividades aprobadas de Fase 2."],
       ["Mes 5 (feb-2027)", "Fase 2 - cierre de adaptación. Se emite la orden de compra del servidor.", "Autorizar orden de compra y costos asociados conforme a cotización aprobada."],
       ["Mes 6 (mar-2027)", "Fase 3 - migración e integración; servidor en tránsito de importación.", "Cubrir migración/integración y obligaciones de importación previamente aprobadas."],
       ["Mes 7 (abr-2027)", "Fase 4 - servidor recibido, instalado y aprovisionado. Pago del equipo físico.", "Desembolso principal de hardware contra recepción/condiciones del proveedor e inventario."],
       ["Mes 8 (may-2027)", "Fase 5 - hardening y continuidad.", "Financiar seguridad, continuidad y servicios aprobados."],
       ["Mes 9 (jun-2027)", "Fase 6 - documentación y capacitación.", "Financiar documentación, capacitación y cierre de pendientes aprobados."],
       ["Mes 10 (jul-2027)", "Fase 7 - producción y entrega formal.", "Desembolsos finales contra puesta en producción, estabilización y acta formal."]],
      anchos=[2.7, 7.3, 7], cambio_filas=(1,))
p("Nota: se agregaron los meses calendario y el tope de dominio y hosting; se corrigió «Adquisicion». En la versión 5 la columna de meses quedó desalineada con los hitos.", cambio=True, tam=9,
  comentario="Decidir: la reunión del 19-sep acordó adelantar la gestión de compra del servidor a enero de 2027 (mes 4). Esta tabla sigue con la orden de compra en el mes 5. Christian (Anexo F) y Víctor (Anexo D) confirman la fecha.")
p("8.5. Víctor Cuenca y el INVERSIONISTA no estarán obligados a efectuar aportes adicionales por sobrecostos, ampliaciones o desviaciones que excedan el máximo expresamente comprometido, salvo aprobación escrita. Ningún espacio en blanco o estimación presupuestaria generará una obligación automática.",
  comentario="Decidir: la reunión del 19-sep acordó que, si el capital de trabajo no alcanza para gastos o remuneraciones posteriores a la entrega, los cinco accionistas aportan en partes iguales. Esta cláusula no lo contradice para los sobrecostos del proyecto, pero conviene dejar escrito el caso del capital de trabajo.")
p("8.6. Todo cambio de alcance, arquitectura crítica, plazo o presupuesto deberá constar en acta de control de cambios con justificación, impacto, responsables y aprobación. Ningún socio podrá modificar unilateralmente los anexos.")

# ------------------------------------------------------------------ 9
h("9. INFRAESTRUCTURA, SEGURIDAD Y CONTINUIDAD")
p("La infraestructura y la seguridad constituyen componentes esenciales del proyecto, con igual exigencia de documentación y aceptación que el desarrollo. Como mínimo deberán producirse y mantenerse los siguientes entregables, cuyo plan y fechas constan en el Anexo F:", cambio=False)
INFRA = [["INF-01", "Diagrama físico y lógico e inventario actualizado", "Christian"],
         ["INF-02", "Baseline de sistemas operativos, hypervisor/servicios y configuraciones", "Christian"],
         ["NET-01", "Segmentación y matriz de flujos de red", "Christian"],
         ["SEC-01", "Políticas de firewall y mínimo privilegio", "Christian"],
         ["SEC-02", "Acceso administrativo seguro: VPN/MFA/SSH/RDP restringido según aplique", "Christian"],
         ["SEC-03", "Hardening de infraestructura y aplicación", "Christian + Desarrollo según capa"],
         ["SEC-04", "TLS y certificados: inventario, vigencia y renovación", "Christian"],
         ["SEC-05", "Gestión segura de credenciales, claves y secretos", "Christian + Desarrollo según capa"],
         ["MON-01", "Monitoreo de infraestructura, BD y aplicación", "Compartido según capa"],
         ["MON-02", "Alertamiento y escalamiento", "Christian"],
         ["LOG-01", "Logging y auditoría de infraestructura, aplicación y tenant", "Compartido"],
         ["BKP-01", "Política de respaldo, retención, cifrado y copia externa", "Christian"],
         ["BKP-02", "Pruebas de restauración con RTO/RPO observado", "Christian + Desarrollo"],
         ["DR-01", "Plan de continuidad y recuperación ante desastres", "Christian"],
         ["VUL-01", "Gestión de vulnerabilidades, parchado y remediación", "Christian + Desarrollo"],
         ["CAP-01", "Capacidad, rendimiento y umbrales", "Christian + Desarrollo"],
         ["DOC-01", "Runbook operativo", "Christian"]]
tabla([["ID", "Entregable mínimo", "Responsable primario"]] + INFRA, anchos=[2, 10, 5])
p("Las credenciales, contraseñas, claves privadas, API keys, tokens, secretos JWT y credenciales de base de datos no deberán almacenarse en repositorios de código. Se administrarán mediante mecanismos separados de custodia, acceso y rotación.")

# ------------------------------------------------------------------ 10
h("10. APORTES, VALORACIÓN Y PARTICIPACIÓN")
p("10.1. Los socios reconocen que existen aportes de naturaleza distinta y que deberán identificarse y valorarse separadamente antes de su capitalización definitiva: software existente; desarrollo futuro; numerario; trabajo técnico y de gestión; derecho de uso de espacio físico; y cualquier otro activo o servicio.")
p("10.2. El supuesto avance del 60% del software no equivale al 60% del valor de TECNIFIN S.A.S. ni determina por sí solo participación accionaria. Su valoración dependerá de activos efectivamente existentes y transferibles, titularidad acreditada, funcionalidad verificable, calidad, deuda técnica, documentación, dependencias y trabajo pendiente.")
p("10.3. El mismo trabajo o activo no podrá contabilizarse simultáneamente como aporte accionario y como servicio íntegramente pagado. Los pagos de desarrollo durante la ejecución deberán separarse de la valoración del software preexistente.")
p("10.4. El aporte máximo inicialmente comprometido por Víctor Cuenca será de hasta USD 20.000 en numerario, sujeto a desembolsos efectivamente realizados y documentados. Los gastos personales que realice fuera de ese monto deberán registrarse separadamente para determinar si corresponden a reembolso, préstamo, aporte adicional u otro tratamiento aprobado. La programación de este aporte se realizará conforme a los hitos de la cláusula 8.4.1 y no se presumirá que Víctor Cuenca debe cubrir con su aporte las obligaciones que correspondan al quinto inversionista o a otras fuentes de financiamiento no materializadas.")
p("10.5. El trabajo especializado de Christian Cuenca será registrado desde el primer día mediante horas y entregables. Para los diez (10) meses se mantiene una valoración referencial total de USD 24.970. De dicho valor, USD 20.000 constituirán la referencia económica de su aporte técnico societario, sujeto a la estructuración jurídica, contable y tributaria que corresponda; y USD 4.970 corresponderán a trabajo especializado adicional remunerable. La forma y calendario de pago de estos USD 4.970 deberá constar en el flujo de caja del Anexo D y vincularse a trabajo efectivamente realizado y documentado, sin confundirse con dividendos ni con el aporte accionario. Las horas y entregables reales deberán registrarse periódicamente.")
tabla([["Función", "Horas ref.", "Tarifa ref.", "Valor ref."],
       ["Jefatura de Proyecto", "300", "USD 13/h", "USD 3.900"],
       ["Infraestructura tecnológica", "360", "USD 27/h", "USD 9.720"],
       ["Seguridad de la información", "300", "USD 32/h", "USD 9.600"],
       ["Administración técnica", "70", "USD 25/h", "USD 1.750"],
       ["TOTAL", "1.030", "", "USD 24.970"]], anchos=[7, 3, 3, 4])
p("Esta valoración es referencial para negociación societaria y registro de aportes; no constituye por sí misma obligación de pago, salario, factura ni determinación tributaria. Las horas y entregables reales deberán registrarse periódicamente.",
  comentario="Decidir: este párrafo dice «no constituye obligación de pago» justo después de declarar USD 4.970 remunerables. Propuesta: limitarlo a los USD 20.000 del aporte accionario.")
p("10.6. Como referencia de negociación, el trabajo histórico preexistente declarado por Franklin Lechón se valora en USD 15.900 y el de Jorge Vladimir Tuquinga Tituaña en USD 20.000, mientras que el trabajo pendiente de los diez meses se compensa hasta por USD 5.250 para cada uno, a razón de USD 700 mensuales durante los meses 1 a 5 y USD 350 mensuales durante los meses 6 a 10. La valoración histórica no se considerará definitiva para fines societarios hasta que se encuentre respaldada por evidencia verificable, incluyendo, según corresponda, historial de repositorios y commits, versiones, módulos, documentación, tickets, pruebas, autoría y titularidad de los activos transferibles. La compensación futura deberá registrarse separadamente del aporte histórico y no podrá generar doble contabilización.",
  comentario="Decidir: las reuniones del 17 y 19-sep acordaron igualar el aporte histórico en USD 20.000 para cada desarrollador (Franklin completa USD 4.100 con trabajo adicional, p. ej. la página web) y dos cuadros de USD 25.250. Además el Anexo H reparte 20 % a cada socio, que supone aportes iguales. Si se confirma, el cuadro 10.6-A pasa a USD 25.250.")
tabla([["Cuadro 10.6-A. Franklin Lechón", "Horas / período ref.", "Tarifa ref.", "Valor ref."],
       ["A. Trabajo histórico ya realizado. Aporte en trabajo, no pagado en efectivo", "", "", ""],
       ["Análisis funcional y diseño del modelo de datos (normativa SEPS, Catálogo Único de Cuentas)", "100", "USD 25/h", "USD 2.500"],
       ["Desarrollo backend: socios, cuentas, caja, créditos, plazo fijo y contabilidad", "240", "USD 25/h", "USD 6.000"],
       ["Desarrollo frontend: interfaces de operación y administración", "120", "USD 25/h", "USD 3.000"],
       ["Reportería regulatoria SEPS/UAF y proceso de cartera", "80", "USD 30/h", "USD 2.400"],
       ["Pruebas, documentación y despliegue de demostración", "80", "USD 25/h", "USD 2.000"],
       ["Subtotal A. Aporte en trabajo histórico", "620", "", "USD 15.900"],
       ["B. Trabajo pendiente del proyecto (40%). Compensación prorrateada en efectivo", "", "", ""],
       ["Fases 1 y 2: arquitectura multi-tenant y adaptación de módulos", "Meses 1 a 5", "USD 700/mes", "USD 3.500"],
       ["Fases 3 a 7: migración, hardening de aplicación, documentación, capacitación y producción", "Meses 6 a 10", "USD 350/mes", "USD 1.750"],
       ["Subtotal B. Compensación prorrateada", "10 meses", "", "USD 5.250"],
       ["TOTAL", "", "", "USD 21.150"]], anchos=[9, 2.8, 2.4, 2.8])
tabla([["Cuadro 10.6-B. Jorge Vladimir Tuquinga Tituaña", "Horas / período ref.", "Tarifa ref.", "Valor ref."],
       ["A. Trabajo histórico ya realizado. Aporte en trabajo, no pagado en efectivo", "", "", ""],
       ["Análisis funcional y diseño del modelo de datos (normativa SEPS, Catálogo Único de Cuentas)", "120", "USD 25/h", "USD 3.000"],
       ["Desarrollo backend: socios, cuentas, caja, créditos, plazo fijo y contabilidad", "300", "USD 25/h", "USD 7.500"],
       ["Desarrollo frontend: interfaces de operación y administración", "160", "USD 25/h", "USD 4.000"],
       ["Reportería regulatoria SEPS/UAF y proceso de cartera", "100", "USD 30/h", "USD 3.000"],
       ["Pruebas, documentación y despliegue de demostración", "100", "USD 25/h", "USD 2.500"],
       ["Subtotal A. Aporte en trabajo histórico", "780", "", "USD 20.000"],
       ["B. Trabajo pendiente del proyecto (40%). Compensación prorrateada en efectivo", "", "", ""],
       ["Fases 1 y 2: arquitectura multi-tenant y adaptación de módulos", "Meses 1 a 5", "USD 700/mes", "USD 3.500"],
       ["Fases 3 a 7: migración, hardening de aplicación, documentación, capacitación y producción", "Meses 6 a 10", "USD 350/mes", "USD 1.750"],
       ["Subtotal B. Compensación prorrateada", "10 meses", "", "USD 5.250"],
       ["TOTAL", "", "", "USD 25.250"]], anchos=[9, 2.8, 2.4, 2.8])
p("10.6.1. Compensación y aceptación de entregables. La compensación mensual por trabajo pendiente se devengará contra la ejecución y evidencia del trabajo correspondiente al período y guardará relación con los hitos y entregables aplicables. Las observaciones menores que no impidan la utilización del entregable no suspenderán injustificadamente el pago; sin embargo, la falta material de cumplimiento o la ausencia de evidencia permitirá diferir el desembolso correspondiente hasta su subsanación. Esta disposición prevalece sobre cualquier redacción que establezca que el pago es independiente de la aceptación técnica.",
  comentario="Decidir (el punto más sensible para los desarrolladores): el 17-sep se aclaró que el pago mensual de 700/350 es independiente del trámite de actas. Propuesta: definir «falta material», fijar 10 días hábiles para observar por escrito y que, sin observación en ese plazo, el pago no se difiera; desempate por los socios.")
p("10.6.2. Terminación anticipada del desarrollo. Como excepción a lo previsto en la cláusula 10.6.1 respecto del devengo mensual; si los socios desarrolladores concluyen antes del mes 10 la totalidad del alcance técnico y funcional a su cargo, y dicho cierre anticipado cuenta con evidencia, pruebas y aceptación formal de los entregables aplicables, conservarán el derecho a percibir la compensación total originalmente pactada para los diez (10) meses. Los pagos pendientes continuarán efectuándose conforme al calendario mensual originalmente acordado hasta el mes 10, salvo que los socios aprueben por escrito su pago anticipado. La terminación anticipada no reducirá la compensación pactada ni generará obligación de extender artificialmente actividades, y los desarrolladores deberán atender durante el período restante las correcciones atribuibles a sus entregables y el soporte de estabilización previsto en el alcance, sin pagos adicionales por dichas correcciones.")
p("10.7. Participación accionaria y cierre de valoración. La participación accionaria definitiva no se fijará únicamente por una cifra preliminar de horas o por el porcentaje declarado de avance del software. Antes de incorporarla al instrumento societario deberán: (i) completarse el inventario y la valoración de los aportes; (ii) verificarse la evidencia de los aportes históricos; (iii) definirse el aporte y condiciones del quinto inversionista; (iv) simularse y dejarse por escrito la distribución porcentual resultante, cuya referencia acordada consta en el cuadro de distribución de acciones del Anexo H; y (v) aprobarse expresamente dicha distribución por todos los socios. Los montos pagados en efectivo como remuneración por servicios no se computarán nuevamente como aporte accionario.", cambio=True,
  comentario="Reunión 21-sep (min 6:18-6:46): la valoración de aportes se reemplaza por un cuadro de distribución de acciones, cinco socios al 20 % cada uno (Anexo H).")

# ------------------------------------------------------------------ 11-16 (sin cambios)
h("11. PROPIEDAD INTELECTUAL, INVENTARIO Y REPOSITORIOS")
p("11.1. Dentro de los cinco (5) días hábiles siguientes a la firma, los desarrolladores entregarán un inventario preliminar de repositorios, módulos, base de datos, documentación, marca y componentes de terceros. Dentro de quince (15) días hábiles entregarán el inventario técnico completo con versiones, estado, evidencia y titularidad. Dentro de veinte (20) días hábiles deberán acreditar la titularidad y los derechos necesarios para la transferencia.")
p("11.2. El inventario deberá incluir, según corresponda: código fuente backend, frontend y móvil; esquemas y scripts de base de datos; procedimientos, funciones y triggers; migraciones/ETL; reportería regulatoria; APIs e integraciones; CI/CD y scripts de despliegue; documentación técnica y funcional; pruebas automatizadas; diseños UI/UX; marca, nombre y logotipo; dominios; librerías propias; componentes de terceros; y datos demo/no productivos.")
p("11.3. Los desarrolladores se obligan a identificar, acreditar y transferir a TECNIFIN S.A.S., mediante los instrumentos jurídicos que correspondan, los derechos patrimoniales transferibles sobre los activos inventariados. La transferencia estará condicionada a la acreditación de titularidad, identificación de componentes de terceros y ausencia de gravámenes o restricciones no reveladas.")
p("11.4. No se considerará cumplido el aporte de la plataforma mediante una declaración genérica de existencia o porcentaje de avance. La aceptación requerirá inventario completo, repositorios accesibles, código compilable/ejecutable conforme a documentación, correspondencia entre versión inventariada y demostrada, dependencias y licencias identificadas, ausencia de credenciales personales embebidas y firma de los instrumentos de cesión aplicables.")
p("11.5. Una vez constituida la sociedad, todo código fuente, documentación técnica, scripts de base de datos, migración y despliegue, configuraciones versionables, pruebas automatizadas, archivos CI/CD y demás activos de desarrollo de la sociedad deberán mantenerse en repositorios institucionales bajo control de TECNIFIN S.A.S. Los repositorios personales podrán utilizarse solo durante la transición y no constituirán el repositorio oficial ni la única fuente de un activo social.")
p("11.6. La transferencia o réplica de repositorios deberá ser íntegra y verificable, preservando cuando sea técnicamente posible historial de commits, ramas, tags, releases y documentación. Deberá identificarse repositorio institucional definitivo, versión/commit de referencia, autoría, titularidad y condición de transferencia.")
p("11.7. Los derechos morales de autor se respetarán conforme a la normativa aplicable. Ningún socio podrá licenciar, ceder o explotar por cuenta propia activos cuya titularidad patrimonial haya sido transferida a la sociedad, salvo autorización societaria expresa.")
h("12. ESPACIO FÍSICO, HARDWARE Y ACTIVOS")
p("12.1. Víctor Cuenca pondrá a disposición de TECNIFIN S.A.S., sin canon de arrendamiento durante un período inicial de seis (6) meses, un espacio físico para alojar la infraestructura tecnológica. Vencido el período de gracia, cualquier canon de arrendamiento o condición de continuidad deberá ser definido y aprobado expresamente; no se presumirá una obligación económica.")
p("12.2. Este derecho de uso no implica transferencia de dominio del inmueble. Deberán documentarse área asignada, plazo, condiciones de acceso, seguridad física, disponibilidad eléctrica, conectividad, ventilación/climatización, adecuaciones, custodia y procedimiento de terminación o traslado.")
p("12.3. La ubicación de activos de TECNIFIN S.A.S. dentro de un inmueble de Víctor Cuenca no modificará la titularidad de dichos activos. Todo equipo comprado con recursos del proyecto deberá ingresar al inventario de la sociedad con identificación, número de serie, costo, responsable de custodia y ubicación.")
p("12.4. Los consumos y servicios directamente atribuibles a la operación de la infraestructura -energía eléctrica incremental, conectividad dedicada, climatización especial u otros costos operativos- no se presumirán incluidos en la cesión gratuita del espacio y deberán ser asumidos, reembolsados o reconocidos por la sociedad conforme al presupuesto aprobado.")
p("12.5. Dominio y hosting. La sociedad contratará en septiembre u octubre de 2026 el dominio y el servicio de alojamiento institucionales, para publicar el sitio web de la compañía y compartir el avance del proyecto, con un presupuesto máximo conjunto de USD 150 incluido en el flujo de caja del Anexo D. Ambos quedarán bajo control institucional conforme a la cláusula 15.2.", cambio=True,
  comentario="Reunión 19-sep (tope USD 150) y 21-sep (min 9:20-9:44): contratar el dominio esta semana o la siguiente para compartir los avances; Jorge pide apoyo con la gestión del pago del hosting.")
h("13. MODELO ECONÓMICO DESDE PRODUCCIÓN")
p("13.1. A partir de la puesta en producción comercial o del mes 11, ninguna remuneración fija a socios por desarrollo, infraestructura, administración, gestión comercial u otros servicios se entenderá garantizada por el solo transcurso del tiempo.")
p("13.2. El órgano competente aprobará periódicamente el presupuesto operativo en función de ingresos efectivamente generados, contratos vigentes, flujo de caja, obligaciones tributarias, costos de operación, reservas y necesidades de reinversión. Los pagos a socios por servicios corresponderán a funciones efectivamente prestadas y serán independientes de dividendos.")
p("13.3. Los ingresos se aplicarán prioritariamente a obligaciones legales y tributarias; costos directos de prestación; infraestructura y seguridad; desarrollo y soporte; administración y comercial; reservas y reinversión; remuneraciones variables adicionales cuando exista capacidad; y finalmente utilidad distribuible conforme a decisión societaria.")
p("13.4. Desarrollo, infraestructura/seguridad, administración y comercialización son funciones necesarias para la operación y ninguna categoría tendrá prioridad económica permanente por defecto. El modelo económico detallado posterior a la producción se definirá por acta al acercarse la puesta en producción.", cambio=True,
  comentario="Reunión 21-sep (min 6:06-6:16): el anexo del modelo económico postproducción (antes G) se retira: «es después». Por eso esta frase remite a un acta futura.")
h("14. INCORPORACIÓN DEL QUINTO INVERSIONISTA")
p("14.1. El quinto inversionista se incorporará mediante la suscripción del instrumento societario correspondiente, adhesión o adenda, según corresponda. Deberán quedar definidos por escrito su identificación completa, aporte, calendario de desembolsos, participación, derechos, obligaciones, tratamiento de dilución y funciones, si las tuviere.")
p("14.2. Hasta que dichas condiciones sean aprobadas, no se presumirá que los socios fundadores deben cubrir proporcionalmente el aporte pendiente del inversionista ni que existe obligación automática de financiamiento adicional.")
p("14.3. Su incorporación no podrá alterar retroactivamente los aportes ya ejecutados ni las obligaciones previamente cumplidas sin consentimiento expreso de los socios afectados.")
h("15. CONFIDENCIALIDAD, CONTROL Y DEPENDENCIA OPERATIVA")
p("15.1. Los socios mantendrán reserva sobre código fuente, arquitectura, credenciales, datos de clientes, información financiera, estrategias comerciales y demás información no pública. La obligación subsistirá por el período que se determine en el instrumento definitivo y por el máximo permitido o exigido para cada categoría de información conforme a la normativa aplicable.")
p("15.2. La compañía deberá controlar institucionalmente repositorios, credenciales, dominios, infraestructura, respaldos y documentación crítica, evitando dependencia exclusiva de una sola persona.")
p("15.3. El acceso a información y sistemas se otorgará bajo necesidad, mínimo privilegio, trazabilidad y revocación oportuna.")
h("16. INCUMPLIMIENTO, RETIRO Y TERMINACIÓN")
p("16.1. El incumplimiento material deberá notificarse por escrito, describiendo la obligación incumplida, evidencia y plazo razonable de subsanación. Para obligaciones técnicas ordinarias se establece como referencia un plazo de diez (10) días hábiles, salvo urgencia, seguridad, continuidad o naturaleza del incumplimiento.")
p("16.2. El retiro de un socio no extinguirá obligaciones pendientes de confidencialidad, propiedad intelectual, entrega de activos, devolución de accesos, documentación o rendición de cuentas.")
p("16.3. Los aportes efectivamente realizados y los trabajos documentados deberán ser conciliados al momento del retiro conforme a su tratamiento societario, contractual, contable y legal. No se presumirá la pérdida automática de trabajo previamente reconocido sin la liquidación correspondiente.")
p("16.4. El incumplimiento grave no subsanado facultará a los demás socios a ejercer las acciones contractuales y legales que correspondan, sin perjuicio de exigir entrega de activos, accesos, documentación y continuidad ordenada.")

# ------------------------------------------------------------------ 17
h("17. ANEXOS VINCULANTES")
p("Los siguientes anexos forman parte integrante e inseparable del contrato, se suscriben junto con él y obligan a los socios desde su firma:", cambio=True)
tabla([["Anexo", "Contenido", "Responsable de su elaboración"],
       ["A", "Alcance, fases, entregables, cronograma y criterios de aceptación", "Jorge Tuquinga y Franklin Lechón"],
       ["B", "Roles y responsabilidades / RACI (incluye funciones de la Gerencia y del Inversionista)", "Jorge Tuquinga y Franklin Lechón; valida Víctor Cuenca"],
       ["C", "Inventario, titularidad, cesión y entrega de propiedad intelectual", "Jorge Tuquinga y Franklin Lechón"],
       ["D", "Flujo de caja de aportes: matriz de aportes, desembolsos y fuentes de financiamiento", "Christian Cuenca, con los datos de Víctor Cuenca"],
       ["E", "Arquitectura multi-tenant, ADR y criterios técnicos de aceptación", "Jorge Tuquinga y Franklin Lechón"],
       ["F", "Plan de Infraestructura Tecnológica y Seguridad de la Información", "Christian Cuenca"],
       ["H", "Cuadro de distribución de acciones", "Christian Cuenca"]],
      anchos=[1.5, 9.5, 6], cambio_filas=(1, 2, 3, 4, 5, 6, 7))
p("El anexo G (modelo económico postproducción) se retira por acuerdo del 21 de septiembre de 2026; ese modelo se definirá por acta conforme a la cláusula 13.4. Cuando un anexo contenga campos pendientes, estos deberán completarse mediante acta antes de que nazca la obligación económica o técnica asociada. Ningún espacio en blanco relativo a porcentajes accionarios, aportes, propiedad intelectual, obligaciones económicas o criterios esenciales de aceptación podrá completarse unilateralmente.", cambio=True,
  comentario="Reparto acordado el 21-sep (min 5:11-7:06): Christian elabora D, F y H; Jorge y Franklin el resto. El flujo de caja va dentro del D (Víctor da los datos).")

h("18. CONTROVERSIAS")
p("Toda controversia derivada de este contrato se procurará resolver primero mediante negociación directa documentada entre los socios. De no alcanzarse acuerdo dentro de treinta (30) días, las partes acudirán a mediación ante un centro de mediación legalmente autorizado en Quito y, de persistir la controversia, a la jurisdicción o mecanismo que se determine en el instrumento definitivo.")
h("19. ACEPTACIÓN Y SUSCRIPCIÓN")
p("Los comparecientes declaran haber leído íntegramente este contrato y sus anexos vinculantes, comprender su contenido y alcance, y se obligan a su cumplimiento. Los porcentajes definitivos, los instrumentos de propiedad intelectual y la validación jurídica correspondiente se formalizarán conforme a las cláusulas 10.7, 11 y 17.", cambio=True)
tabla([["SOCIO", "FIRMA", "C.I."],
       ["Víctor Emilio Cuenca Caraguay - GERENTE", "", "1101475232"],
       ["Franklin Sebastián Lechón Quilo - PRESIDENTE / DESARROLLADOR", "", "1753860053"],
       ["Christian Cuenca - JEFE DE PROYECTO / INFRAESTRUCTURA Y SEGURIDAD", "", "1716860851"],
       ["Jorge Vladimir Tuquinga Tituaña - DESARROLLADOR", "", "1720884012"],
       ["____________________ - INVERSIONISTA", "", "____________"]],
      anchos=[9, 5, 3], cambio_filas=(1, 2, 3, 5))

# ================================================================== ANEXOS
salto(); h("ANEXO A. ALCANCE, FASES, ENTREGABLES, CRONOGRAMA Y CRITERIOS DE ACEPTACIÓN")
p("Responsables: Jorge Tuquinga y Franklin Lechón. Plazo del proyecto: 1 de octubre de 2026 a 31 de julio de 2027 (cláusula 2.3).", cambio=True)
h("A.1 Cronograma de fases", 2)
tabla([["Fase", "Ventana", "Cierre (acta)", "Entregables (cláusulas 7.6 y 9)"],
       ["1 Arquitectura y núcleo multi-tenant", "1-oct-2026 a 30-nov-2026", "30-nov-2026 (hito H1: modelo de datos, 30-oct-2026)", "ARQ-01, ARQ-02, DAT-01, APP-01"],
       ["2 Adaptación de módulos", "1-dic-2026 a 28-feb-2027", "28-feb-2027", "MOD-01, MOD-02, MOD-03, QA-01"],
       ["3 Migración e integración", "1 a 31-mar-2027", "31-mar-2027", "QA-01, MIG-01, INT-01"],
       ["4 Infraestructura física/lógica", "1 a 30-abr-2027", "30-abr-2027", "DEP-01, INF-01, INF-02, NET-01, SEC-01, SEC-02, SEC-04"],
       ["5 Hardening, seguridad, observabilidad y continuidad", "1 a 31-may-2027", "31-may-2027", "APP-02, BKP-03, SEC-03, SEC-05, MON-01, MON-02, LOG-01, BKP-01, BKP-02, DR-01, VUL-01, CAP-01"],
       ["6 Documentación y capacitación", "1 a 30-jun-2027", "30-jun-2027", "DOC-01, DOC-02, CAP-02"],
       ["7 Producción y entrega", "1 a 31-jul-2027", "31-jul-2027", "PRD-01 y acta formal de cierre"]], anchos=[4.2, 3.6, 3.8, 5.4])
p("Fase 3. La plataforma TECNIFIN nace sin datos de ninguna cooperativa. La integración de punta a punta se prueba sobre una base de demostración separada, con datos sintéticos. La migración controlada de datos (MIG-01) aplica cuando una cooperativa cliente traiga datos de su sistema anterior, y se ensaya siempre primero sobre una restauración del respaldo.",
  comentario="Decisión técnica de Jorge del 20-sep (base en blanco). Afecta el sentido de MIG-01 en la cláusula 7.6; confirmar entre los socios.")
h("A.2 Criterios de aceptación comunes", 2)
for x in ["Cada entregable se presenta al Jefe de Proyecto con su evidencia (cláusula 8.1) y se acepta o se observa por acta; nunca por silencio (8.2).",
          "Entregables de software: repositorio institucional con el código en la versión presentada; pruebas automatizadas en verde (en el proyecto, el comando «npm run verificar»); migraciones de base de datos aplicadas sin pendientes ni alteraciones; sin credenciales en el repositorio.",
          "Entregables multi-tenant: prueba de aislamiento entre dos cooperativas sin filtración de datos; toda tabla de negocio con identificador de cooperativa y seguridad por filas forzada.",
          "Entregables contables y de cartera: asientos cuadrados (partida doble), cuentas del Catálogo Único por código numérico y bandas SEPS leídas de datos, no fijas en el código.",
          "Documentos (arquitectura, ADR, manuales): revisados por Desarrollo e Infraestructura/Seguridad cuando toquen ambas capas (cláusula 6.3).",
          "Observaciones: se subsanan en el plazo del entregable o en diez (10) días hábiles (cláusula 8.3)."]:
    p("• " + x)
h("A.3 Estado al 22-sep-2026 (informativo)", 2)
p("Adelantado respecto del cronograma y pendiente de acta: proyecto TECNIFIN creado con reglas de calidad comprobadas por pruebas; borradores de ADR-0001 a ADR-0003 y ARQ-01 (estado «propuesta», pendientes de revisión de Christian Cuenca); modelo de datos multi-tenant construido (40 tablas, seguridad por filas forzada, 54 pruebas automatizadas en verde) y base de demostración.")

salto(); h("ANEXO B. ROLES Y RESPONSABILIDADES / RACI")
p("Responsables: Jorge Tuquinga y Franklin Lechón; Víctor Cuenca valida las funciones de la Gerencia y del Inversionista. R = realiza · A = aprueba / rinde cuentas · C = consultado · I = informado.", cambio=True)
h("B.1 Funciones por socio", 2)
tabla([["Socio", "Función", "Compromisos"],
       ["Víctor Cuenca", "Gerente; aportante en numerario (cl. 2.4, 4.7, 10.4, 12)", "Proveer los datos del flujo de caja de aportes (Anexo D) y actualizarlo; administrar el capital de trabajo y los fondos del proyecto con trazabilidad; tratamiento contable de aportes, software y cuentas por pagar durante la constitución; desembolsar su aporte de hasta USD 20.000 según los hitos de la cláusula 8.4.1; poner a disposición el espacio físico (cl. 12); gestionar la incorporación del Inversionista."],
       ["Inversionista", "Aportante en numerario (cl. 2.5, 4.8, 14)", "Desembolsar su aporte de hasta USD 20.000 según los hitos de la cláusula 8.4.1; apoyar la socialización del producto, el relacionamiento institucional y la comercialización. Funciones a confirmar con su incorporación."],
       ["Christian Cuenca", "Jefe de Proyecto; Infraestructura y Seguridad (cl. 4.4, 4.5)", "Cronograma, riesgos, cambios y actas; revisión expresa de ADR críticas; entregables de la cláusula 9 y Anexo F; elaborar los Anexos D, F y H."],
       ["Franklin Lechón", "Presidente; Desarrollo (cl. 4.6, 7)", "Entregables de la cláusula 7.6 en conjunto con Jorge; frontend, pruebas y soporte a usuarios; reportería SEPS; elaborar los Anexos A, B, C y E con Jorge."],
       ["Jorge Tuquinga", "Desarrollo (cl. 4.6, 7)", "Entregables de la cláusula 7.6 en conjunto con Franklin; backend, base de datos y arquitectura multi-tenant; inventario de propiedad intelectual; elaborar los Anexos A, B, C y E con Franklin."]],
      anchos=[3.2, 4.3, 9.5], cambio_filas=(1, 2))
h("B.2 Matriz RACI de actividades clave", 2)
tabla([["Actividad", "Víctor", "Inversionista", "Christian", "Franklin", "Jorge"],
       ["Flujo de caja de aportes (Anexo D)", "R", "C", "A", "I", "I"],
       ["Desembolsos de aportes en numerario", "R/A", "R", "C", "I", "I"],
       ["Capital de trabajo y tratamiento contable", "R/A", "I", "C", "I", "I"],
       ["Cronograma, riesgos y control de cambios", "I", "I", "R/A", "C", "C"],
       ["Actas de aceptación de entregables", "I", "I", "R/A", "C", "C"],
       ["Arquitectura multi-tenant y ADR", "I", "I", "A (capa infra)", "R", "R"],
       ["Desarrollo de módulos y pruebas", "I", "I", "C", "R/A", "R/A"],
       ["Infraestructura, seguridad y respaldos", "I", "I", "R/A", "C", "C"],
       ["Inventario y cesión de propiedad intelectual", "I", "I", "C", "R", "R"],
       ["Dominio, hosting y sitio web", "C (pago)", "I", "A", "R (sitio)", "R (dominio)"],
       ["Socialización y comercialización", "R", "R", "C", "C", "C"],
       ["Capacitación funcional", "I", "I", "C", "R", "R"]], anchos=[5.6, 1.8, 2.3, 2.6, 2.3, 2.4], tam=9)

salto(); h("ANEXO C. INVENTARIO, TITULARIDAD, CESIÓN Y ENTREGA DE PROPIEDAD INTELECTUAL")
p("Responsables: Jorge Tuquinga y Franklin Lechón. Inventario preliminar; se completa en los plazos de la cláusula 11.1 (5, 15 y 20 días hábiles desde la firma).", cambio=True)
tabla([["Activo", "Descripción", "Ubicación actual", "Titular a acreditar", "Estado"],
       ["Plataforma TECNIFIN", "Plataforma multi-tenant sobre PostgreSQL (modelo de datos, migraciones, pruebas, documentación, ADR)", "Repositorio privado github.com/JorgeVladimir/TECNIFIN", "Jorge Tuquinga / Franklin Lechón", "En desarrollo"],
       ["Sistema financiero existente (GUTT SYSTEM)", "Sistema mono-cooperativa de referencia funcional: backend, frontend, reportería SEPS, scripts de base de datos", "Repositorio privado github.com/JorgeVladimir/GUTT_SYSTEM", "Jorge Tuquinga / Franklin Lechón", "En producción (referencia)"],
       ["Aplicación móvil del socio", "Aplicación móvil y su diseño de 17 pantallas", "Repositorio GUTT_SYSTEM_MOVIL", "Jorge Tuquinga / Franklin Lechón", "Diseño hecho; construcción pendiente"],
       ["Marca, nombre y logotipo", "TECNIFIN S.A.S. / GUTT COMPANY S.A.S.", "Archivos de diseño", "A definir", "Pendiente de registro"],
       ["Dominio", "tecnifin.com (a contratar, cl. 12.5)", "—", "TECNIFIN S.A.S.", "Pendiente"],
       ["Datos de demostración", "Base de demostración con datos sintéticos", "Repositorio TECNIFIN (scripts)", "TECNIFIN S.A.S.", "Hecho"],
       ["Componentes de terceros", "Librerías de código abierto (p. ej. PostgreSQL, Node.js, librería pg)", "Dependencias declaradas en cada repositorio", "Terceros (licencias)", "Por listar con sus licencias"]],
      anchos=[3.2, 4.8, 3.8, 3, 2.2], tam=8.5)
p("Cesión: una vez acreditada la titularidad (cl. 11.3), los desarrolladores firman el instrumento de cesión de derechos patrimoniales a TECNIFIN S.A.S. y los repositorios se trasladan a la cuenta institucional de la sociedad preservando su historial (cl. 11.5 y 11.6).")

salto(); h("ANEXO D. FLUJO DE CAJA DE APORTES: MATRIZ DE APORTES, DESEMBOLSOS Y FUENTES DE FINANCIAMIENTO")
p("Responsable: Christian Cuenca, con los datos que provee Víctor Cuenca. ESQUEMA DE TRABAJO: los montos ya acordados se incluyen como referencia; Christian y Víctor completan las fechas y las entradas por aportante.", cambio=True)
tabla([["Mes", "Desarrolladores (Jorge + Franklin)", "Christian (10.5)", "Dominio y hosting", "Servidor e infraestructura", "Otros", "Aporte Víctor", "Aporte Inversionista"],
       ["Oct-2026", "1.400", "≈ 497", "≤ 150", "", "", "", ""],
       ["Nov-2026", "1.400", "≈ 497", "", "", "", "", ""],
       ["Dic-2026", "1.400", "≈ 497", "", "", "≈ 100 (reunión navideña)", "", ""],
       ["Ene-2027", "1.400", "≈ 497", "", "gestión de compra", "", "", ""],
       ["Feb-2027", "1.400", "≈ 497", "", "orden de compra", "", "", ""],
       ["Mar-2027", "700", "≈ 497", "", "", "", "", ""],
       ["Abr-2027", "700", "≈ 497", "", "10.569,45 (proforma)", "", "", ""],
       ["May-2027", "700", "≈ 497", "", "", "", "", ""],
       ["Jun-2027", "700", "≈ 497", "", "", "", "", ""],
       ["Jul-2027", "700", "≈ 497", "", "", "", "", ""],
       ["TOTAL (USD)", "10.500", "4.970", "≤ 150", "10.569,45", "≈ 100", "≤ 20.000", "≤ 20.000"]],
      anchos=[1.8, 2.3, 1.8, 1.8, 2.5, 2.4, 2.2, 2.2], tam=8)
p("Referencia: egresos cuantificados ≈ USD 26.289,45 frente a un numerario comprometido de hasta USD 40.000 (Víctor e Inversionista). Por conciliar: infraestructura 10.569,45 (proforma) frente a 10.595,45 (citado en la reunión del 19-sep); estimación del servidor básico con traslado desde EE. UU.; internet redundante; trámites de constitución.")

salto(); h("ANEXO E. ARQUITECTURA MULTI-TENANT, ADR Y CRITERIOS TÉCNICOS DE ACEPTACIÓN")
p("Responsables: Jorge Tuquinga y Franklin Lechón. Revisión expresa de Christian Cuenca en lo que toque aislamiento, seguridad, respaldo o capacidad (cl. 6.3).", cambio=True)
tabla([["Documento", "Contenido", "Estado"],
       ["ARQ-01", "Documento de arquitectura multi-tenant con los apartados de la cláusula 6.1", "Borrador"],
       ["ADR-0001", "Identificación del tenant (cooperativa) y numeración por cooperativa desde 1", "Propuesta"],
       ["ADR-0002", "Aislamiento de datos: una base, un esquema, identificador de cooperativa y seguridad por filas forzada; decisión provisional sobre la fijación del tenant", "Propuesta: requiere revisión expresa de Christian"],
       ["ADR-0003", "Modelo y nomenclatura: tipos (dinero con dos decimales, fechas), catálogos por cooperativa, imágenes dentro de la base", "Propuesta"]], anchos=[2.4, 10.2, 4.4])
h("Criterios técnicos de aceptación", 2)
for x in ["Toda tabla de negocio lleva el identificador de cooperativa, con seguridad por filas habilitada y forzada; una prueba recorre el catálogo de la base y falla si alguna tabla queda sin protección.",
          "Una cooperativa no puede leer, modificar ni referenciar datos de otra (prueba automatizada con dos cooperativas).",
          "El rol con el que se conecta la aplicación no es dueño de las tablas ni puede saltarse la seguridad por filas.",
          "Importes monetarios con dos decimales y sin valores no numéricos; partida doble exigida por la base de datos.",
          "El plan de cuentas es propio de cada cooperativa; ninguna comparte catálogo con otra.",
          "La base de producción nace sin datos; los datos de ejemplo viven solo en la base de demostración.",
          "Riesgos abiertos que debe revisar Christian: tiempo de restauración de una cooperativa sin afectar a las demás, volumen y retención, cifrado y custodia de claves, y el riesgo residual de la fijación del tenant (ADR-0002)."]:
    p("• " + x)

salto(); h("ANEXO F. PLAN DE INFRAESTRUCTURA TECNOLÓGICA Y SEGURIDAD DE LA INFORMACIÓN")
p("Responsable: Christian Cuenca. ESQUEMA DE TRABAJO: se listan los entregables de la cláusula 9 para que Christian complete la fase, la fecha objetivo y la evidencia de cada uno.", cambio=True)
tabla([["ID", "Entregable", "Fase", "Fecha objetivo", "Evidencia"]] + [[a, b, "", "", ""] for a, b, _ in INFRA], anchos=[1.8, 8, 1.4, 2.4, 3.4], tam=8.5)

salto(); h("ANEXO H. CUADRO DE DISTRIBUCIÓN DE ACCIONES")
p("Responsable: Christian Cuenca. Acordado en la reunión del 21-sep-2026: cinco socios, veinte por ciento (20 %) cada uno, sujeto a la cláusula 10.7.", cambio=True,
  comentario="Si se mantiene el 20 % para todos, el aporte histórico de Franklin (USD 15.900 en el cuadro 10.6-A) debe igualarse a USD 20.000 como se acordó el 17 y 19-sep; si no, hay que decidir cómo se refleja la diferencia.")
tabla([["Socio", "Cédula", "Participación"],
       ["Víctor Emilio Cuenca Caraguay", "1101475232", "20 %"],
       ["Franklin Sebastián Lechón Quilo", "1753860053", "20 %"],
       ["Christian Cuenca", "1716860851", "20 %"],
       ["Jorge Vladimir Tuquinga Tituaña", "1720884012", "20 %"],
       ["____________________ (Inversionista)", "____________", "20 %"],
       ["TOTAL", "", "100 %"]], anchos=[8, 4, 4])

doc.save(OUT)
print("guardado:", OUT)
